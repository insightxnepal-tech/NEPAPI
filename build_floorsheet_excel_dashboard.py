#!/usr/bin/env python3
"""
Build an interactive NEPSE floorsheet Excel dashboard.

Creates a multi-sheet .xlsx with:
  - Dashboard KPIs + charts
  - Daily / monthly market pulse
  - Top stocks & broker net flow
  - Stock×Day fact table (PivotTable / Slicer ready)
  - Whale / block trades
  - How-to sheet for Excel interactivity

Usage:
  python build_floorsheet_excel_dashboard.py
  python build_floorsheet_excel_dashboard.py --out ~/Downloads/floorsheet/NEPSE_Floorsheet_Dashboard.xlsx
"""

from __future__ import annotations

import argparse
import glob
import os
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo


# ── theme ────────────────────────────────────────────────────────────────────
NAVY = "0B3D5C"
TEAL = "0E7C7B"
GOLD = "C4A35A"
CREAM = "F7F4EF"
GREEN = "1B7F4E"
RED = "B42318"
WHITE = "FFFFFF"
GRAY = "6B7280"
LIGHT = "E8EEF2"

thin = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB"),
)


def downloads_dir() -> Path:
    preferred = Path("/Users/sanishtamang/Downloads/floorsheet")
    if preferred.parent.exists() or preferred.exists():
        preferred.mkdir(parents=True, exist_ok=True)
        return preferred
    fallback = Path.home() / "Downloads" / "floorsheet"
    fallback.mkdir(parents=True, exist_ok=True)
    return fallback


def load_floorsheets(data_dir: str = ".") -> pd.DataFrame:
    files = sorted(glob.glob(os.path.join(data_dir, "floorsheet_2026-*.csv")))
    files = [f for f in files if "dividend" not in f]
    if not files:
        raise FileNotFoundError("No floorsheet_2026-*.csv files found.")

    usecols = [
        "contractId",
        "stockSymbol",
        "buyerMemberId",
        "sellerMemberId",
        "contractQuantity",
        "contractRate",
        "contractAmount",
        "businessDate",
        "buyerBrokerName",
        "sellerBrokerName",
    ]
    frames = []
    for path in files:
        header = pd.read_csv(path, nrows=0).columns.tolist()
        cols = [c for c in usecols if c in header]
        frames.append(pd.read_csv(path, usecols=cols))

    data = pd.concat(frames, ignore_index=True)
    for col in ("contractAmount", "contractQuantity", "contractRate"):
        if col in data.columns:
            data[col] = pd.to_numeric(data[col], errors="coerce")
    data["businessDate"] = pd.to_datetime(data["businessDate"], errors="coerce")
    if "contractId" in data.columns:
        data = data.drop_duplicates(subset=["contractId", "businessDate"], keep="first")
    for col in ("buyerBrokerName", "sellerBrokerName"):
        if col in data.columns:
            data[col] = data[col].fillna("").astype(str).replace({"nan": ""})
    data = data.dropna(subset=["businessDate", "stockSymbol", "contractAmount"])
    return data.reset_index(drop=True)


def build_aggregates(data: pd.DataFrame) -> dict[str, pd.DataFrame]:
    daily = (
        data.groupby("businessDate", as_index=False)
        .agg(
            Turnover=("contractAmount", "sum"),
            Volume=("contractQuantity", "sum"),
            Trades=("contractId", "count"),
            Scrips=("stockSymbol", "nunique"),
            AvgRate=("contractRate", "mean"),
        )
        .sort_values("businessDate")
    )
    daily["TurnoverCr"] = daily["Turnover"] / 1e7
    daily["Date"] = daily["businessDate"].dt.strftime("%Y-%m-%d")

    monthly = (
        data.assign(Month=data["businessDate"].dt.to_period("M").astype(str))
        .groupby("Month", as_index=False)
        .agg(
            Turnover=("contractAmount", "sum"),
            Volume=("contractQuantity", "sum"),
            Trades=("contractId", "count"),
            TradingDays=("businessDate", "nunique"),
        )
    )
    monthly["TurnoverCr"] = monthly["Turnover"] / 1e7

    stocks = (
        data.groupby("stockSymbol", as_index=False)
        .agg(
            Turnover=("contractAmount", "sum"),
            Volume=("contractQuantity", "sum"),
            Trades=("contractId", "count"),
            Days=("businessDate", "nunique"),
            AvgRate=("contractRate", "mean"),
            LastRate=("contractRate", "last"),
        )
        .sort_values("Turnover", ascending=False)
    )
    stocks["TurnoverCr"] = stocks["Turnover"] / 1e7
    stocks["ChangePct"] = (
        (stocks["LastRate"] - stocks["AvgRate"]) / stocks["AvgRate"] * 100
    ).round(2)

    stock_daily = (
        data.groupby(["businessDate", "stockSymbol"], as_index=False)
        .agg(
            Open=("contractRate", "first"),
            High=("contractRate", "max"),
            Low=("contractRate", "min"),
            Close=("contractRate", "last"),
            Volume=("contractQuantity", "sum"),
            Turnover=("contractAmount", "sum"),
            Trades=("contractId", "count"),
        )
        .sort_values(["businessDate", "Turnover"], ascending=[True, False])
    )
    stock_daily["Date"] = stock_daily["businessDate"].dt.strftime("%Y-%m-%d")
    stock_daily["TurnoverCr"] = stock_daily["Turnover"] / 1e7

    buy = (
        data.groupby("buyerBrokerName", as_index=False)["contractAmount"]
        .sum()
        .rename(columns={"buyerBrokerName": "Broker", "contractAmount": "Bought"})
    )
    sell = (
        data.groupby("sellerBrokerName", as_index=False)["contractAmount"]
        .sum()
        .rename(columns={"sellerBrokerName": "Broker", "contractAmount": "Sold"})
    )
    brokers = buy.merge(sell, on="Broker", how="outer").fillna(0)
    brokers = brokers[brokers["Broker"].astype(str).str.len() > 0]
    brokers["Net"] = brokers["Bought"] - brokers["Sold"]
    brokers["NetCr"] = brokers["Net"] / 1e7
    brokers["BoughtCr"] = brokers["Bought"] / 1e7
    brokers["SoldCr"] = brokers["Sold"] / 1e7
    brokers = brokers.sort_values("Net", ascending=False)

    whale = (
        data.nlargest(200, "contractAmount")[
            [
                "businessDate",
                "stockSymbol",
                "contractQuantity",
                "contractRate",
                "contractAmount",
                "buyerBrokerName",
                "sellerBrokerName",
                "buyerMemberId",
                "sellerMemberId",
            ]
        ]
        .copy()
    )
    whale["Date"] = whale["businessDate"].dt.strftime("%Y-%m-%d")
    whale["AmountCr"] = whale["contractAmount"] / 1e7
    whale = whale.rename(
        columns={
            "stockSymbol": "Symbol",
            "contractQuantity": "Qty",
            "contractRate": "Rate",
            "buyerBrokerName": "BuyerBroker",
            "sellerBrokerName": "SellerBroker",
            "buyerMemberId": "BuyerId",
            "sellerMemberId": "SellerId",
        }
    )

    # Heat: late 10 sessions vs early 10
    dates = sorted(data["businessDate"].unique())
    if len(dates) >= 20:
        early = data[data["businessDate"].isin(dates[:10])]
        late = data[data["businessDate"].isin(dates[-10:])]
        e = early.groupby("stockSymbol")["contractAmount"].sum()
        l = late.groupby("stockSymbol")["contractAmount"].sum()
        heat = pd.DataFrame({"Early10d": e, "Late10d": l}).fillna(0)
        heat = heat[(heat["Early10d"] > 5e6) & (heat["Late10d"] > 5e6)]
        heat["Multiple"] = (heat["Late10d"] / heat["Early10d"]).round(2)
        heat["EarlyCr"] = heat["Early10d"] / 1e7
        heat["LateCr"] = heat["Late10d"] / 1e7
        heat = heat.reset_index().rename(columns={"stockSymbol": "Symbol"})
        heat = heat.sort_values("Multiple", ascending=False)
    else:
        heat = pd.DataFrame(columns=["Symbol", "EarlyCr", "LateCr", "Multiple"])

    return {
        "daily": daily,
        "monthly": monthly,
        "stocks": stocks,
        "stock_daily": stock_daily,
        "brokers": brokers,
        "whale": whale,
        "heat": heat,
    }


def style_header(ws, row: int, start: int, end: int) -> None:
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(bold=True, color=WHITE, name="Calibri", size=11)
    for col in range(start, end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin


def autosize(ws, min_width: int = 10, max_width: int = 28) -> None:
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        length = 0
        for cell in col[:80]:
            if cell.value is None:
                continue
            length = max(length, len(str(cell.value)))
        ws.column_dimensions[letter].width = max(min_width, min(max_width, length + 2))


def write_table(
    ws,
    df: pd.DataFrame,
    start_row: int,
    start_col: int,
    table_name: str,
    money_cols: set[str] | None = None,
) -> tuple[int, int]:
    money_cols = money_cols or set()
    rows = list(dataframe_to_rows(df, index=False, header=True))
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row, start=start_col):
            cell = ws.cell(row=start_row + r_idx, column=c_idx, value=value)
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
            if r_idx == 0:
                continue
            header = rows[0][c_idx - start_col]
            if header in money_cols and isinstance(value, (int, float)):
                cell.number_format = '#,##0.00'
            elif isinstance(value, float):
                cell.number_format = '#,##0.00'
            elif isinstance(value, int):
                cell.number_format = '#,##0'
    end_row = start_row + len(rows) - 1
    end_col = start_col + len(df.columns) - 1
    style_header(ws, start_row, start_col, end_col)
    ref = (
        f"{get_column_letter(start_col)}{start_row}:"
        f"{get_column_letter(end_col)}{end_row}"
    )
    table = Table(displayName=table_name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)
    return end_row, end_col


def kpi_box(ws, row: int, col: int, title: str, value: str, fill: str) -> None:
    title_cell = ws.cell(row=row, column=col, value=title)
    title_cell.font = Font(name="Calibri", size=10, color=WHITE, bold=True)
    title_cell.fill = PatternFill("solid", fgColor=fill)
    title_cell.alignment = Alignment(horizontal="center")
    val_cell = ws.cell(row=row + 1, column=col, value=value)
    val_cell.font = Font(name="Calibri", size=16, color=NAVY, bold=True)
    val_cell.fill = PatternFill("solid", fgColor=CREAM)
    val_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    for r in (row, row + 1):
        for c in (col, col + 1):
            ws.cell(row=r, column=c).border = thin


def build_workbook(data: pd.DataFrame, aggs: dict[str, pd.DataFrame]) -> Workbook:
    wb = Workbook()

    daily = aggs["daily"]
    monthly = aggs["monthly"]
    stocks = aggs["stocks"]
    stock_daily = aggs["stock_daily"]
    brokers = aggs["brokers"]
    whale = aggs["whale"]
    heat = aggs["heat"]

    total_turnover = float(data["contractAmount"].sum())
    total_volume = float(data["contractQuantity"].sum())
    total_trades = int(len(data))
    n_days = int(data["businessDate"].nunique())
    n_scrips = int(data["stockSymbol"].nunique())
    start = data["businessDate"].min().date()
    end = data["businessDate"].max().date()
    recent3 = daily["Turnover"].tail(3).mean()
    prior = daily["Turnover"].iloc[:-3].mean() if len(daily) > 3 else daily["Turnover"].mean()
    if recent3 > prior * 1.05:
        bias = "BULLISH"
        bias_fill = GREEN
    elif recent3 < prior * 0.95:
        bias = "CAUTIOUS"
        bias_fill = RED
    else:
        bias = "NEUTRAL"
        bias_fill = GOLD

    # ── Dashboard ────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Dashboard"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "NEPSE FLOORSHEET INTERACTIVE DASHBOARD"
    ws["A1"].font = Font(name="Calibri", size=20, bold=True, color=NAVY)
    ws.merge_cells("A1:H1")
    ws["A2"] = (
        f"Period {start} → {end}  ·  Generated {datetime.now():%Y-%m-%d %H:%M}  ·  "
        f"Filter any Excel Table column · Use PivotTables on StockDaily / BrokerFlow"
    )
    ws["A2"].font = Font(name="Calibri", size=10, color=GRAY)
    ws.merge_cells("A2:H2")

    kpi_box(ws, 4, 1, "TOTAL TURNOVER", f"Rs {total_turnover/1e9:.2f} B", NAVY)
    kpi_box(ws, 4, 3, "TRADING DAYS", f"{n_days}", TEAL)
    kpi_box(ws, 4, 5, "CONTRACTS", f"{total_trades:,}", GOLD)
    kpi_box(ws, 4, 7, "MARKET BIAS", bias, bias_fill)

    kpi_box(ws, 7, 1, "AVG DAILY TURNOVER", f"Rs {daily['Turnover'].mean()/1e9:.2f} B", TEAL)
    kpi_box(ws, 7, 3, "TOTAL VOLUME", f"{total_volume/1e6:.1f} M sh", NAVY)
    kpi_box(ws, 7, 5, "UNIQUE SCRIPS", f"{n_scrips}", GOLD)
    kpi_box(
        ws,
        7,
        7,
        "LATEST DAY",
        f"Rs {daily['Turnover'].iloc[-1]/1e9:.2f} B",
        TEAL,
    )

    # Mini daily table for chart
    ws["A10"] = "DAILY TURNOVER (Rs Cr)"
    ws["A10"].font = Font(bold=True, color=NAVY, size=12)
    chart_df = daily[["Date", "TurnoverCr", "Trades", "Scrips"]].copy()
    chart_df.columns = ["Date", "Turnover_Cr", "Trades", "Scrips"]
    end_row, end_col = write_table(ws, chart_df, 11, 1, "DashDaily", {"Turnover_Cr"})

    chart = LineChart()
    chart.title = "Daily Turnover (Rs Crore)"
    chart.style = 10
    chart.y_axis.title = "Rs Cr"
    chart.x_axis.title = None
    chart.height = 10
    chart.width = 18
    data_ref = Reference(ws, min_col=2, min_row=11, max_col=2, max_row=end_row)
    cats = Reference(ws, min_col=1, min_row=12, max_row=end_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    ws.add_chart(chart, "F11")

    # Top 10 stocks snapshot
    top10 = stocks.head(10)[
        ["stockSymbol", "TurnoverCr", "Volume", "Trades", "ChangePct"]
    ].copy()
    top10.columns = ["Symbol", "Turnover_Cr", "Volume", "Trades", "VsAvg_%"]
    ws["A" + str(end_row + 3)] = "TOP 10 STOCKS BY TURNOVER"
    ws["A" + str(end_row + 3)].font = Font(bold=True, color=NAVY, size=12)
    t_start = end_row + 4
    t_end, _ = write_table(
        ws, top10, t_start, 1, "DashTopStocks", {"Turnover_Cr", "VsAvg_%"}
    )
    ws.conditional_formatting.add(
        f"E{t_start+1}:E{t_end}",
        ColorScaleRule(
            start_type="num",
            start_value=-10,
            start_color=RED,
            mid_type="num",
            mid_value=0,
            mid_color=WHITE,
            end_type="num",
            end_value=10,
            end_color=GREEN,
        ),
    )

    bar = BarChart()
    bar.type = "col"
    bar.title = "Top 10 Turnover (Rs Cr)"
    bar.style = 10
    bar.height = 10
    bar.width = 15
    bref = Reference(ws, min_col=2, min_row=t_start, max_row=t_end)
    bcats = Reference(ws, min_col=1, min_row=t_start + 1, max_row=t_end)
    bar.add_data(bref, titles_from_data=True)
    bar.set_categories(bcats)
    bar.shape = 4
    ws.add_chart(bar, "G" + str(t_start))

    autosize(ws)
    ws.row_dimensions[5].height = 28
    ws.row_dimensions[8].height = 28

    # ── DailyPulse ───────────────────────────────────────────────────────
    ws_d = wb.create_sheet("DailyPulse")
    ws_d["A1"] = "Daily Market Pulse — Excel Table (filter / sort / slicer-ready)"
    ws_d["A1"].font = Font(bold=True, size=14, color=NAVY)
    ddf = daily[
        ["Date", "Turnover", "TurnoverCr", "Volume", "Trades", "Scrips", "AvgRate"]
    ].copy()
    ddf.columns = [
        "Date",
        "Turnover_Rs",
        "Turnover_Cr",
        "Volume",
        "Trades",
        "Scrips",
        "AvgRate",
    ]
    write_table(
        ws_d,
        ddf,
        3,
        1,
        "DailyPulse",
        {"Turnover_Rs", "Turnover_Cr", "AvgRate"},
    )
    ws_d.conditional_formatting.add(
        f"C4:C{3+len(ddf)}",
        DataBarRule(start_type="min", end_type="max", color=TEAL),
    )
    autosize(ws_d)

    # ── Monthly ──────────────────────────────────────────────────────────
    ws_m = wb.create_sheet("Monthly")
    ws_m["A1"] = "Monthly Turnover Summary"
    ws_m["A1"].font = Font(bold=True, size=14, color=NAVY)
    mdf = monthly.copy()
    mdf.columns = [
        "Month",
        "Turnover_Rs",
        "Volume",
        "Trades",
        "TradingDays",
        "Turnover_Cr",
    ]
    end_m, _ = write_table(
        ws_m, mdf, 3, 1, "MonthlyPulse", {"Turnover_Rs", "Turnover_Cr"}
    )
    mchart = BarChart()
    mchart.type = "col"
    mchart.title = "Monthly Turnover (Rs Cr)"
    mchart.style = 10
    mchart.height = 10
    mchart.width = 12
    mref = Reference(ws_m, min_col=6, min_row=3, max_row=end_m)
    mcats = Reference(ws_m, min_col=1, min_row=4, max_row=end_m)
    mchart.add_data(mref, titles_from_data=True)
    mchart.set_categories(mcats)
    ws_m.add_chart(mchart, "H3")
    autosize(ws_m)

    # ── TopStocks ────────────────────────────────────────────────────────
    ws_s = wb.create_sheet("TopStocks")
    ws_s["A1"] = "All Stocks by Period Turnover — filter Symbol / sort Change%"
    ws_s["A1"].font = Font(bold=True, size=14, color=NAVY)
    sdf = stocks[
        [
            "stockSymbol",
            "Turnover",
            "TurnoverCr",
            "Volume",
            "Trades",
            "Days",
            "AvgRate",
            "LastRate",
            "ChangePct",
        ]
    ].copy()
    sdf.columns = [
        "Symbol",
        "Turnover_Rs",
        "Turnover_Cr",
        "Volume",
        "Trades",
        "Days",
        "AvgRate",
        "LastRate",
        "VsAvg_%",
    ]
    end_s, _ = write_table(
        ws_s,
        sdf,
        3,
        1,
        "TopStocks",
        {"Turnover_Rs", "Turnover_Cr", "AvgRate", "LastRate", "VsAvg_%"},
    )
    ws_s.conditional_formatting.add(
        f"I4:I{end_s}",
        ColorScaleRule(
            start_type="percentile",
            start_value=10,
            start_color=RED,
            mid_type="percentile",
            mid_value=50,
            mid_color=WHITE,
            end_type="percentile",
            end_value=90,
            end_color=GREEN,
        ),
    )
    autosize(ws_s)

    # ── BrokerFlow ───────────────────────────────────────────────────────
    ws_b = wb.create_sheet("BrokerFlow")
    ws_b["A1"] = "Broker Net Flow — positive = accumulation, negative = distribution"
    ws_b["A1"].font = Font(bold=True, size=14, color=NAVY)
    bdf = brokers[
        ["Broker", "Bought", "Sold", "Net", "BoughtCr", "SoldCr", "NetCr"]
    ].copy()
    bdf.columns = [
        "Broker",
        "Bought_Rs",
        "Sold_Rs",
        "Net_Rs",
        "Bought_Cr",
        "Sold_Cr",
        "Net_Cr",
    ]
    end_b, _ = write_table(
        ws_b,
        bdf,
        3,
        1,
        "BrokerFlow",
        {"Bought_Rs", "Sold_Rs", "Net_Rs", "Bought_Cr", "Sold_Cr", "Net_Cr"},
    )
    ws_b.conditional_formatting.add(
        f"G4:G{end_b}",
        ColorScaleRule(
            start_type="min",
            start_color=RED,
            mid_type="num",
            mid_value=0,
            mid_color=WHITE,
            end_type="max",
            end_color=GREEN,
        ),
    )
    # Top accumulators / distributors chart data
    top_acc = bdf.head(10)
    ws_b["I3"] = "Top Accumulators (Net Cr)"
    ws_b["I3"].font = Font(bold=True, color=GREEN)
    for i, row in enumerate(top_acc.itertuples(index=False), start=4):
        ws_b.cell(row=i, column=9, value=row.Broker[:32])
        ws_b.cell(row=i, column=10, value=row.Net_Cr)
    bchart = BarChart()
    bchart.type = "bar"
    bchart.title = "Top 10 Net Accumulators (Rs Cr)"
    bchart.style = 10
    bchart.height = 12
    bchart.width = 14
    bchart.add_data(Reference(ws_b, min_col=10, min_row=3, max_row=13), titles_from_data=False)
    bchart.set_categories(Reference(ws_b, min_col=9, min_row=4, max_row=13))
    bchart.dataLabels = DataLabelList()
    ws_b.add_chart(bchart, "I16")
    autosize(ws_b)

    # ── StockDaily (Pivot source) ────────────────────────────────────────
    ws_sd = wb.create_sheet("StockDaily")
    ws_sd["A1"] = (
        "FACT TABLE — Insert → PivotTable for interactive Symbol/Date analysis. "
        "Recommended: Rows=Symbol, Columns=Date, Values=Turnover_Cr / Volume"
    )
    ws_sd["A1"].font = Font(bold=True, size=12, color=NAVY)
    # Keep workbook size manageable: top 80 symbols by turnover × all days
    top_syms = set(stocks.head(80)["stockSymbol"])
    sdfact = stock_daily[stock_daily["stockSymbol"].isin(top_syms)].copy()
    sdfact = sdfact[
        [
            "Date",
            "stockSymbol",
            "Open",
            "High",
            "Low",
            "Close",
            "Volume",
            "Turnover",
            "TurnoverCr",
            "Trades",
        ]
    ]
    sdfact.columns = [
        "Date",
        "Symbol",
        "Open",
        "High",
        "Low",
        "Close",
        "Volume",
        "Turnover_Rs",
        "Turnover_Cr",
        "Trades",
    ]
    write_table(
        ws_sd,
        sdfact,
        3,
        1,
        "StockDaily",
        {"Open", "High", "Low", "Close", "Turnover_Rs", "Turnover_Cr"},
    )
    autosize(ws_sd)

    # ── Momentum ─────────────────────────────────────────────────────────
    ws_h = wb.create_sheet("Momentum")
    ws_h["A1"] = "Turnover Multiple — Late 10 sessions vs Early 10 sessions"
    ws_h["A1"].font = Font(bold=True, size=14, color=NAVY)
    if len(heat):
        hdf = heat[["Symbol", "EarlyCr", "LateCr", "Multiple"]].copy()
        hdf.columns = ["Symbol", "Early10d_Cr", "Late10d_Cr", "Multiple"]
        end_h, _ = write_table(
            ws_h, hdf, 3, 1, "MomentumHeat", {"Early10d_Cr", "Late10d_Cr", "Multiple"}
        )
        ws_h.conditional_formatting.add(
            f"D4:D{end_h}",
            ColorScaleRule(
                start_type="num",
                start_value=0.2,
                start_color=RED,
                mid_type="num",
                mid_value=1,
                mid_color=WHITE,
                end_type="num",
                end_value=5,
                end_color=GREEN,
            ),
        )
    autosize(ws_h)

    # ── WhaleTrades ──────────────────────────────────────────────────────
    ws_w = wb.create_sheet("WhaleTrades")
    ws_w["A1"] = "Largest 200 contracts by amount — filter by Symbol / Date / Broker"
    ws_w["A1"].font = Font(bold=True, size=14, color=NAVY)
    wdf = whale[
        [
            "Date",
            "Symbol",
            "Qty",
            "Rate",
            "contractAmount",
            "AmountCr",
            "BuyerId",
            "SellerId",
            "BuyerBroker",
            "SellerBroker",
        ]
    ].copy()
    wdf.columns = [
        "Date",
        "Symbol",
        "Qty",
        "Rate",
        "Amount_Rs",
        "Amount_Cr",
        "BuyerId",
        "SellerId",
        "BuyerBroker",
        "SellerBroker",
    ]
    write_table(
        ws_w,
        wdf,
        3,
        1,
        "WhaleTrades",
        {"Rate", "Amount_Rs", "Amount_Cr"},
    )
    autosize(ws_w)

    # ── HowTo ────────────────────────────────────────────────────────────
    ws_i = wb.create_sheet("HowTo_Interactive")
    ws_i.sheet_view.showGridLines = False
    ws_i["A1"] = "How to use this workbook interactively in Excel"
    ws_i["A1"].font = Font(size=16, bold=True, color=NAVY)
    tips = [
        "",
        "1. Every data sheet uses an Excel Table (filter arrows on headers).",
        "2. Click any column filter to slice by Date, Symbol, Broker, etc.",
        "3. Create a PivotTable from StockDaily:",
        "      Select any cell in the StockDaily table → Insert → PivotTable",
        "      Rows: Symbol   Columns: Date   Values: Sum of Turnover_Cr",
        "4. Add Slicers: PivotTable Analyze → Insert Slicer → Symbol / Date",
        "5. BrokerFlow: Pivot on Broker with Values = Net_Cr for smart-money view.",
        "6. Momentum sheet: Multiple > 2 = heating up; Multiple < 0.5 = cooling.",
        "7. WhaleTrades: sort Amount_Cr descending to spot block deals.",
        "8. Dashboard charts update when you refresh source CSVs and re-run:",
        "      python build_floorsheet_excel_dashboard.py",
        "",
        "Sheets:",
        "  Dashboard   — KPI cards + daily/top-stock charts",
        "  DailyPulse  — one row per trading day",
        "  Monthly     — month rollup",
        "  TopStocks   — full scrip ranking",
        "  BrokerFlow  — net buy/sell by broker",
        "  StockDaily  — OHLCV fact table (best for Pivot + Slicers)",
        "  Momentum    — late vs early turnover multiples",
        "  WhaleTrades — largest contracts",
    ]
    for i, line in enumerate(tips, start=2):
        ws_i.cell(row=i, column=1, value=line).font = Font(
            name="Calibri", size=12, color=NAVY if line.startswith(" ") or line.endswith(":") else "1F2937"
        )
    ws_i.column_dimensions["A"].width = 90

    # Freeze panes on data sheets
    for name in ("DailyPulse", "TopStocks", "BrokerFlow", "StockDaily", "Momentum", "WhaleTrades"):
        wb[name].freeze_panes = "A4"

    return wb


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data-dir",
        default=".",
        help="Directory containing floorsheet_YYYY-MM-DD.csv files",
    )
    parser.add_argument(
        "--out",
        default=None,
        help="Output .xlsx path (default: ~/Downloads/floorsheet/NEPSE_Floorsheet_Dashboard.xlsx)",
    )
    args = parser.parse_args()

    print("Loading floorsheets …")
    data = load_floorsheets(args.data_dir)
    print(
        f"  {len(data):,} contracts · {data['businessDate'].nunique()} days · "
        f"{data['stockSymbol'].nunique()} scrips"
    )
    print("Aggregating …")
    aggs = build_aggregates(data)
    print("Building workbook …")
    wb = build_workbook(data, aggs)

    out = Path(args.out) if args.out else downloads_dir() / "NEPSE_Floorsheet_Dashboard.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)

    # Also mirror next to repo for git / CI
    mirror = Path(args.data_dir) / "NEPSE_Floorsheet_Dashboard.xlsx"
    if mirror.resolve() != out.resolve():
        wb.save(mirror)
        print(f"Mirrored → {mirror}")

    print(f"Dashboard saved → {out}")
    print(f"Size: {out.stat().st_size / 1e6:.1f} MB")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
